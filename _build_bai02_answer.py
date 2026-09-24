import zipfile
from copy import deepcopy
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from lxml import etree

SRC = Path(r"C:\Users\GHC\Downloads\Bai tap 02_MaSinhVien_HoTenSinhVien.docx")
TMP = Path(r"C:\Users\GHC\AppData\Local\Temp\aims-bai02-task")
OUT = Path(r"C:\Dev\Learning\VIII.HUST\Phân tích và xây dựng phần mềm\01_Word\Bai tap 02_202490077_BachMinhQuang.docx")
IMG_DIR = TMP / "diagrams"
TMP.mkdir(parents=True, exist_ok=True)
IMG_DIR.mkdir(parents=True, exist_ok=True)
FONT_REG = r"C:\Windows\Fonts\arial.ttf"
FONT_BOLD = r"C:\Windows\Fonts\arialbd.ttf"

cases = [
{
"code":"UC001","name":"Place Order","actors":"Customer",
"brief":"Khách hàng đặt các sản phẩm trong giỏ; AIMS kiểm tra tồn kho, thu thập thông tin giao hàng, tính hóa đơn và ghi nhận đơn sau khi thanh toán thành công.",
"pre":"Giỏ hàng có ít nhất một sản phẩm với số lượng dương. Khách hàng đang ở màn hình giỏ hàng; thông tin giao hàng có thể được cung cấp hợp lệ.",
"steps":[
("Customer","Mở giỏ hàng và yêu cầu xem các sản phẩm dự định mua."),
("AIMS Software","Kiểm tra số lượng tồn kho hiện tại của từng sản phẩm trong giỏ."),
("AIMS Software","Hiển thị danh sách sản phẩm, số lượng, đơn giá và thành tiền."),
("Customer","Chọn sản phẩm, điều chỉnh số lượng cần mua và yêu cầu đặt hàng."),
("AIMS Software","Hiển thị biểu mẫu nhập thông tin giao hàng."),
("Customer","Nhập tên, số điện thoại, địa chỉ nhận hàng và phương thức giao hàng."),
("AIMS Software","Kiểm tra trường bắt buộc, định dạng liên hệ và khả năng giao tới địa chỉ."),
("AIMS Software","Tính phí giao hàng và tổng tiền dựa trên sản phẩm đã chọn."),
("AIMS Software","Hiển thị hóa đơn tạm gồm chi tiết sản phẩm, phí giao và tổng thanh toán."),
("Customer","Kiểm tra hóa đơn, xác nhận đặt đơn và yêu cầu thanh toán."),
("AIMS Software","Tạo đơn chờ thanh toán, gắn mã đơn và giữ số lượng hàng trong thời gian thanh toán."),
("AIMS Software","Gọi UC002 Pay Order với mã đơn và tổng tiền phải trả."),
("UC002 Pay Order","Trả kết quả thanh toán thành công cùng mã giao dịch cho AIMS."),
("AIMS Software","Ghi nhận đơn đã thanh toán, chuyển trạng thái sang chờ xử lý và xóa sản phẩm đã đặt khỏi giỏ."),
("AIMS Software","Gửi email xác nhận có mã đơn, thông tin giao hàng và hóa đơn; hiển thị thông báo đặt hàng thành công.")
],
"alts":[
("2a","AIMS Software","Nếu tồn kho không đủ, thông báo số lượng khả dụng và yêu cầu khách hàng cập nhật hoặc bỏ sản phẩm; quay lại bước 1."),
("7a","AIMS Software","Nếu thông tin giao hàng thiếu, sai định dạng hoặc ngoài vùng phục vụ, đánh dấu trường lỗi và yêu cầu khách hàng sửa; quay lại bước 6."),
("12a","AIMS Software","Nếu thanh toán bị từ chối hoặc chưa có kết quả, giữ giỏ và đơn ở trạng thái chưa thanh toán; cho phép thử lại hoặc hủy, không ghi nhận đơn đã trả tiền.")
],
"post":"Khi thành công, đơn hàng và giao dịch được lưu, trạng thái đơn là đã thanh toán/chờ xử lý và giỏ được cập nhật. Nếu thanh toán chưa thành công, AIMS không xác nhận đơn đã trả tiền và giữ dữ liệu để khách hàng tiếp tục.",
"prompt":"Prompt đã dùng: Hãy đặc tả UC Place Order cho AIMS theo mẫu, gồm tác nhân, tiền điều kiện, 15 bước thành công, ba luồng thay thế, hậu điều kiện và dữ liệu vào/ra.",
"eval":"Đánh giá, điều chỉnh: Luồng được tách rõ giữa tạo đơn chờ thanh toán và xác nhận đơn sau khi UC002 thành công. Khi thiếu hàng, dữ liệu giao hàng sai hoặc thanh toán lỗi, hệ thống không xóa giỏ và không báo đặt hàng thành công.",
"io":[
("Sản phẩm, số lượng trong giỏ","Vào","Mã sản phẩm và số lượng nguyên dương; AIMS kiểm tra tồn kho tại thời điểm đặt."),
("Tên, điện thoại, địa chỉ nhận","Vào","Bắt buộc; kiểm tra định dạng và vùng giao hàng."),
("Phương thức giao hàng","Vào","Chọn giao thường; lựa chọn giao nhanh được xử lý tại UC003."),
("Xác nhận đặt hàng","Vào","Thao tác xác nhận của Customer sau khi kiểm tra hóa đơn."),
("Hóa đơn tạm","Ra","Chi tiết hàng, số lượng, đơn giá, phí giao hàng và tổng tiền."),
("Mã đơn và trạng thái","Ra","Mã đơn; trạng thái chờ thanh toán hoặc đã thanh toán/chờ xử lý."),
("Email xác nhận","Ra","Mã đơn, sản phẩm, tổng tiền và thông tin giao hàng.")
],
"diagram":{
"nodes":[
("s",0,0,"","start"),("cart",1,0,"Xem giỏ hàng","action"),
("stock",2,1,"Kiểm tra tồn kho","action"),("stock_q",3,1,"Đủ hàng?","decision"),
("stock_no",4,0,"Báo thiếu hàng; cập nhật giỏ","action"),("cart_show",4,1,"Hiển thị giỏ hàng","action"),
("delivery",5,0,"Chọn hàng và yêu cầu đặt","action"),("form",6,1,"Hiển thị biểu mẫu giao hàng","action"),
("address",7,0,"Nhập thông tin nhận hàng","action"),("validate",8,1,"Kiểm tra thông tin, tính phí","action"),
("invoice",9,1,"Hiển thị hóa đơn","action"),("confirm",10,0,"Xác nhận đơn và yêu cầu trả tiền","action"),
("pay",11,1,"Gọi UC002 Pay Order","action"),("paid_q",12,1,"Thanh toán thành công?","decision"),
("pay_no",13,0,"Thông báo lỗi; thử lại hoặc hủy","action"),("commit",13,1,"Lưu đơn, xóa giỏ, gửi xác nhận","action"),
("e",14,0,"Customer nhận xác nhận","end")],
"edges":[
("s","cart",""),("cart","stock",""),("stock","stock_q",""),("stock_q","stock_no","Không"),
("stock_no","cart","Sửa giỏ"),("stock_q","cart_show","Có"),("cart_show","delivery",""),
("delivery","form",""),("form","address",""),("address","validate",""),("validate","invoice",""),
("invoice","confirm",""),("confirm","pay",""),("pay","paid_q",""),("paid_q","pay_no","Không"),
("pay_no","confirm","Thử lại"),("paid_q","commit","Có"),("commit","e","")]
}},
{
"code":"UC002","name":"Pay Order","actors":"Customer; Interbank",
"brief":"Customer thanh toán tổng tiền của đơn bằng thẻ; AIMS gửi yêu cầu tới Interbank, ghi nhận kết quả và trả trạng thái cho UC gọi.",
"pre":"AIMS đã có đơn chờ thanh toán và tính được tổng tiền, mã đơn, loại tiền. Kết nối bảo mật tới Interbank đang sẵn sàng.",
"steps":[
("AIMS Software","Hiển thị màn hình thanh toán với mã đơn, tổng tiền và loại tiền."),
("Customer","Nhập thông tin chủ thẻ và thông tin thẻ cần thiết."),
("Customer","Kiểm tra số tiền rồi xác nhận thanh toán."),
("AIMS Software","Kiểm tra trường bắt buộc, định dạng số thẻ, ngày hết hạn và mã bảo mật."),
("AIMS Software","Tạo yêu cầu thanh toán an toàn gồm mã đơn, số tiền và dữ liệu thẻ được bảo vệ."),
("AIMS Software","Gửi yêu cầu xác thực/giao dịch tới Interbank qua kết nối bảo mật."),
("Interbank","Xác thực thẻ, kiểm tra hạn mức và quyết định chấp thuận hoặc từ chối."),
("Interbank","Trả kết quả, mã phản hồi và mã giao dịch cho AIMS."),
("AIMS Software","Đối chiếu mã đơn, số tiền, loại tiền và mã giao dịch để tránh ghi nhận sai hoặc trùng."),
("AIMS Software","Lưu số tiền, thời điểm, mã giao dịch, trạng thái và thông tin thẻ đã che một phần."),
("AIMS Software","Nếu được chấp thuận, cập nhật trạng thái thanh toán là PAID."),
("AIMS Software","Trả kết quả thành công cho UC Place Order hoặc UC Place Rush Order đang gọi."),
("AIMS Software","Hiển thị xác nhận thanh toán cùng mã giao dịch và đơn hàng."),
("Customer","Nhận thông báo kết quả thanh toán trên AIMS.")
],
"alts":[
("4a","AIMS Software","Nếu thiếu hoặc sai định dạng thông tin thẻ, không gửi yêu cầu tới Interbank; hiển thị trường cần sửa và quay lại bước 2."),
("7a","Interbank","Nếu thẻ không hợp lệ, hết hạn hoặc không đủ hạn mức, trả kết quả từ chối; AIMS ghi nhận thất bại, thông báo và cho phép dùng thẻ khác hoặc hủy."),
("7b","AIMS Software","Nếu hết thời gian chờ hoặc mất kết nối nên chưa biết kết quả, ghi trạng thái PENDING và tra cứu giao dịch trước khi gửi lại để tránh trừ tiền hai lần.")
],
"post":"Khi được chấp thuận, giao dịch được lưu và thanh toán chuyển sang PAID. Khi bị từ chối hoặc chưa rõ kết quả, đơn không được đánh dấu đã thanh toán; kết quả lỗi/chờ xử lý được ghi nhận. AIMS chỉ lưu thông tin thẻ đã che một phần, không lưu mã bảo mật.",
"prompt":"Prompt đã dùng: Hãy đặc tả UC Pay Order của AIMS theo mẫu với Customer và Interbank, bao gồm xác thực thẻ, phản hồi ngân hàng, lỗi thẻ, hết hạn mức và timeout.",
"eval":"Đánh giá, điều chỉnh: Luồng thanh toán được tách khỏi UC Place Order. Phản hồi AI ban đầu chưa phân biệt giao dịch bị từ chối với timeout chưa rõ kết quả; đặc tả bổ sung trạng thái PENDING và bước tra cứu trước khi gửi lại, đồng thời không lưu mã bảo mật thẻ.",
"io":[
("Mã đơn, số tiền, loại tiền","Vào","Do UC đặt hàng gửi; số tiền phải bằng tổng tiền đã xác nhận."),
("Tên chủ thẻ và dữ liệu thẻ/token","Vào","Customer cung cấp; truyền qua kết nối bảo mật, che một phần khi lưu."),
("Xác nhận thanh toán","Vào","Customer xác nhận giao dịch với số tiền hiển thị."),
("Yêu cầu giao dịch","Ra","AIMS gửi mã đơn, số tiền, loại tiền và dữ liệu thanh toán được bảo vệ tới Interbank."),
("Kết quả, mã giao dịch","Vào từ Interbank","Chấp thuận, từ chối hoặc chưa xác định; kèm mã phản hồi/mã giao dịch nếu có."),
("Trạng thái thanh toán","Ra","PAID, FAILED hoặc PENDING; trả cho UC gọi và hiển thị cho Customer."),
("Thông tin giao dịch","Ra","Số tiền, thời điểm, mã giao dịch và thẻ đã che một phần; không lưu CVV/mã bảo mật.")
],
"diagram":{
"nodes":[
("s",0,0,"","start"),("screen",1,1,"Hiển thị số tiền và đơn hàng","action"),
("card",2,0,"Nhập thẻ và xác nhận","action"),("format",3,1,"Kiểm tra định dạng","action"),
("format_q",4,1,"Dữ liệu hợp lệ?","decision"),("format_no",5,0,"Sửa thông tin thẻ","action"),
("request",5,1,"Gửi yêu cầu an toàn","action"),("process",6,2,"Interbank xử lý thẻ","action"),
("approved_q",7,2,"Được chấp thuận?","decision"),("declined",8,0,"Nhận lỗi; đổi thẻ hoặc hủy","action"),
("record",8,1,"Lưu mã giao dịch và kết quả","action"),("paid",9,1,"Cập nhật PAID","action"),
("return",10,1,"Trả kết quả cho UC gọi","action"),("e",11,0,"Customer nhận kết quả","end")],
"edges":[
("s","screen",""),("screen","card",""),("card","format",""),("format","format_q",""),
("format_q","format_no","Không"),("format_no","card","Sửa"),("format_q","request","Có"),
("request","process",""),("process","approved_q",""),("approved_q","declined","Không"),
("declined","card","Thử lại"),("approved_q","record","Có"),("record","paid",""),
("paid","return",""),("return","e","")]
}},
{
"code":"UC003","name":"Place Rush Order","actors":"Customer",
"brief":"Customer chọn giao hàng nhanh; AIMS kiểm tra khả năng phục vụ, tính phụ phí và thời gian dự kiến rồi thanh toán, lưu đơn ưu tiên.",
"pre":"Giỏ hàng có ít nhất một sản phẩm và Customer đang đặt hàng. AIMS có địa chỉ nhận hàng để kiểm tra vùng giao nhanh.",
"steps":[
("Customer","Chọn tùy chọn giao hàng nhanh trong quá trình checkout."),
("AIMS Software","Kiểm tra sản phẩm trong giỏ có hỗ trợ giao nhanh hay không."),
("AIMS Software","Kiểm tra địa chỉ thuộc vùng phục vụ và thời điểm đặt còn trong giờ nhận đơn nhanh."),
("AIMS Software","Tính phí giao nhanh theo địa chỉ và sản phẩm."),
("AIMS Software","Ước tính khoảng thời gian giao hàng nhanh có thể đáp ứng."),
("AIMS Software","Hiển thị phí giao nhanh, thời gian dự kiến và hóa đơn cập nhật."),
("Customer","Xem thông tin, xác nhận phí và thông tin giao hàng."),
("AIMS Software","Kiểm tra lại dữ liệu giao hàng và điều kiện áp dụng trước khi chốt."),
("AIMS Software","Cập nhật đơn tạm với loại giao hàng nhanh và phí tương ứng."),
("AIMS Software","Tính và hiển thị tổng tiền cuối cùng cần thanh toán."),
("Customer","Xác nhận đặt đơn giao nhanh và yêu cầu thanh toán."),
("AIMS Software","Gọi UC002 Pay Order với mã đơn và tổng tiền cuối cùng."),
("UC002 Pay Order","Trả kết quả thanh toán cho AIMS."),
("AIMS Software","Nếu thanh toán thành công, ghi nhận đơn đã trả tiền và gắn mức ưu tiên giao nhanh."),
("AIMS Software","Giữ hàng, xóa sản phẩm đã đặt khỏi giỏ và gửi email xác nhận phí, thời gian giao dự kiến.")
],
"alts":[
("2a","AIMS Software","Nếu sản phẩm, địa chỉ hoặc thời điểm không hỗ trợ giao nhanh, thông báo lý do và cho Customer chuyển sang giao thường; tiếp tục UC001."),
("6a","Customer","Nếu Customer không đồng ý phụ phí hoặc thời gian dự kiến, bỏ lựa chọn giao nhanh và tiếp tục UC001 với giao thường."),
("12a","AIMS Software","Nếu thanh toán bị từ chối hoặc timeout, không xác nhận đơn giao nhanh đã thanh toán; giữ giỏ và cho thử lại hoặc hủy.")
],
"post":"Nếu thành công, đơn được ghi nhận với lựa chọn giao nhanh, phụ phí và mức ưu tiên; thông báo xác nhận được gửi cho Customer. Nếu không đủ điều kiện hoặc thanh toán lỗi, đơn giao nhanh chưa được xác nhận và có thể tiếp tục bằng giao thường.",
"prompt":"Prompt đã dùng: Hãy đặc tả UC Place Rush Order cho AIMS theo mẫu, tập trung vào kiểm tra vùng giao, điều kiện sản phẩm, phụ phí, thời gian dự kiến và thanh toán.",
"eval":"Đánh giá, điều chỉnh: Phản hồi AI ban đầu coi giao nhanh là lựa chọn luôn khả dụng. Nội dung được bổ sung điều kiện theo sản phẩm, địa chỉ và giờ nhận đơn; nếu không khả dụng hoặc Customer không đồng ý phí, luồng quay về giao thường thay vì tạo đơn sai cam kết.",
"io":[
("Giỏ hàng và số lượng","Vào","Sản phẩm cần đặt; AIMS phải kiểm tra lại tồn kho."),
("Địa chỉ và thông tin liên hệ","Vào","Dùng để kiểm tra vùng giao nhanh và tính phí."),
("Lựa chọn giao nhanh","Vào","Customer chọn giao nhanh khi checkout; có thể bỏ chọn để về UC001."),
("Phí và thời gian dự kiến","Ra","AIMS trả phụ phí và khoảng thời gian giao dự kiến trước khi xác nhận."),
("Tổng tiền cuối cùng","Ra","Tổng tiền hàng cộng phí giao nhanh."),
("Mã đơn và mức ưu tiên","Ra","Được ghi nhận sau khi UC002 trả về thanh toán thành công."),
("Email xác nhận","Ra","Mã đơn, lựa chọn giao nhanh, phí và thời gian dự kiến.")
],
"diagram":{
"nodes":[
("s",0,0,"","start"),("choose",1,0,"Chọn giao hàng nhanh","action"),
("check",2,1,"Kiểm tra sản phẩm, địa chỉ, giờ nhận","action"),
("eligible_q",3,1,"Có thể giao nhanh?","decision"),("unavailable",4,0,"Chuyển giao thường; tiếp tục UC001","action"),
("fee",4,1,"Tính phụ phí và thời gian dự kiến","action"),("quote",5,1,"Hiển thị hóa đơn cập nhật","action"),
("confirm",6,0,"Xác nhận lựa chọn và phí","action"),("accept_q",7,1,"Customer đồng ý?","decision"),
("reject",8,0,"Bỏ tùy chọn giao nhanh","action"),("update",8,1,"Cập nhật loại giao và tổng tiền","action"),
("pay",9,1,"Gọi UC002 Pay Order","action"),("save",10,1,"Lưu đơn ưu tiên sau khi trả tiền","action"),
("e",11,0,"Gửi email xác nhận","end")],
"edges":[
("s","choose",""),("choose","check",""),("check","eligible_q",""),("eligible_q","unavailable","Không"),
("eligible_q","fee","Có"),("fee","quote",""),("quote","confirm",""),("confirm","accept_q",""),
("accept_q","reject","Không"),("reject","unavailable",""),("accept_q","update","Có"),
("update","pay",""),("pay","save",""),("save","e","")]
}},
{
"code":"UC004","name":"Cancel Order","actors":"Customer",
"brief":"Customer yêu cầu hủy đơn chưa bàn giao cho đơn vị giao hàng; AIMS kiểm tra quyền và trạng thái, xử lý hoàn tiền nếu đơn đã thanh toán.",
"pre":"Đơn tồn tại và thuộc Customer đang yêu cầu. Đơn chưa được bàn giao/gửi đi; nếu đã thanh toán thì có giao dịch gốc để UC005 kiểm tra hoàn tiền.",
"steps":[
("Customer","Mở lịch sử đơn hàng và chọn đơn muốn hủy."),
("Customer","Chọn chức năng hủy đơn và nhập lý do."),
("AIMS Software","Tải thông tin đơn, chủ sở hữu, trạng thái giao hàng và thanh toán."),
("AIMS Software","Xác thực Customer là chủ đơn hoặc có quyền hủy."),
("AIMS Software","Kiểm tra đơn chưa được bàn giao cho đơn vị giao hàng và còn trong điều kiện hủy."),
("AIMS Software","Xác định đơn đã thanh toán hay chưa và số tiền đã thu."),
("AIMS Software","Hiển thị sản phẩm, trạng thái, lý do và khoản dự kiến hoàn nếu đã trả tiền."),
("Customer","Kiểm tra thông tin và xác nhận yêu cầu hủy."),
("AIMS Software","Ghi nhận yêu cầu hủy có mã tham chiếu để không xử lý lặp."),
("AIMS Software","Nếu đơn đã thanh toán, gọi UC005 Refund Order với giao dịch gốc; nếu chưa trả tiền thì bỏ qua bước hoàn tiền."),
("UC005 Refund Order","Trả kết quả hoàn tiền thành công hoặc đang chờ xử lý (nếu có)."),
("AIMS Software","Ngừng các bước xử lý/giao hàng tiếp theo và chuyển đơn sang trạng thái CANCELLED."),
("AIMS Software","Cập nhật trạng thái hoàn tiền thành công hoặc REFUND_PENDING theo kết quả UC005."),
("AIMS Software","Giải phóng số lượng đã giữ nếu đơn chưa được bàn giao."),
("AIMS Software","Gửi email và hiển thị kết quả hủy cùng trạng thái hoàn tiền cho Customer.")
],
"alts":[
("3a","AIMS Software","Nếu không tìm thấy đơn hoặc Customer không sở hữu đơn, từ chối yêu cầu và không tiết lộ thông tin đơn của người khác; kết thúc UC."),
("5a","AIMS Software","Nếu đơn đã được bàn giao/gửi đi hoặc không còn đủ điều kiện hủy, thông báo lý do và giữ nguyên trạng thái đơn; kết thúc UC."),
("10a","AIMS Software","Nếu UC005 chưa hoàn tiền do Interbank lỗi/timeout, vẫn ghi nhận hủy với trạng thái REFUND_PENDING, thông báo Customer và cho phép tra cứu/hỗ trợ; không gửi lặp giao dịch.")
],
"post":"Đơn đủ điều kiện được chuyển CANCELLED và dừng xử lý giao hàng. Nếu chưa thanh toán thì không tạo giao dịch hoàn tiền; nếu đã thanh toán thì trạng thái hoàn tiền liên kết với UC005 là thành công hoặc đang chờ xử lý.",
"prompt":"Prompt đã dùng: Hãy đặc tả UC Cancel Order cho AIMS theo mẫu; kiểm tra chủ sở hữu, trạng thái giao hàng, xác nhận hủy và liên kết UC Refund Order khi đơn đã thanh toán.",
"eval":"Đánh giá, điều chỉnh: Phản hồi AI ban đầu giả định mọi đơn đều có thể hủy. Đặc tả bổ sung kiểm tra quyền sở hữu và thời điểm trước khi bàn giao hàng; đơn đã thanh toán gọi UC005, còn trạng thái hoàn tiền được giữ riêng để không báo đã hoàn khi Interbank chưa xác nhận.",
"io":[
("Mã đơn, phiên Customer","Vào","AIMS dùng để tìm đơn và xác thực chủ sở hữu."),
("Lý do hủy","Vào","Lý do do Customer nhập; lưu cùng lịch sử xử lý."),
("Xác nhận hủy","Vào","Customer xác nhận sau khi xem trạng thái và khoản dự kiến hoàn."),
("Kết quả kiểm tra điều kiện","Ra","Đủ điều kiện, không đủ điều kiện hoặc đơn đã bàn giao."),
("Trạng thái hủy đơn","Ra","CANCELLED khi yêu cầu hợp lệ được ghi nhận."),
("Trạng thái/số tiền hoàn","Ra","Không áp dụng nếu chưa trả tiền; nếu đã trả tiền lấy kết quả và số tiền từ UC005."),
("Email kết quả","Ra","Thông báo hủy và trạng thái hoàn tiền hiện tại.")
],
"diagram":{
"nodes":[
("s",0,0,"","start"),("select",1,0,"Chọn đơn, nhập lý do hủy","action"),
("check",2,1,"Kiểm tra chủ đơn và trạng thái giao","action"),("eligible_q",3,1,"Đủ điều kiện hủy?","decision"),
("reject",4,0,"Thông báo không thể hủy","end"),("impact",4,1,"Hiển thị tác động và tiền dự kiến hoàn","action"),
("confirm",5,0,"Xác nhận hủy","action"),("paid_q",6,1,"Đơn đã thanh toán?","decision"),
("refund",7,1,"Gọi UC005 Refund Order","action"),("skip",7,0,"Bỏ qua hoàn tiền","action"),
("update",8,1,"Hủy đơn, dừng giao, giải phóng hàng","action"),
("notify",9,1,"Lưu trạng thái hoàn tiền và gửi email","action"),
("e",10,0,"Customer nhận kết quả","end")],
"edges":[
("s","select",""),("select","check",""),("check","eligible_q",""),("eligible_q","reject","Không"),
("eligible_q","impact","Có"),("impact","confirm",""),("confirm","paid_q",""),
("paid_q","refund","Có"),("paid_q","skip","Chưa"),("refund","update",""),("skip","update",""),
("update","notify",""),("notify","e","")]
}},
{
"code":"UC005","name":"Refund Order","actors":"Customer; Interbank",
"brief":"AIMS kiểm tra một giao dịch đã thu tiền có đủ điều kiện hoàn hay không, gửi yêu cầu tới Interbank và cập nhật kết quả hoàn tiền cho đơn.",
"pre":"Đơn có giao dịch thanh toán thành công và đã bị hủy hoặc được xác nhận đủ điều kiện hoàn. Số tiền đề nghị không vượt số đã thu và chưa hoàn.",
"steps":[
("Customer","Mở đơn đã hủy/đủ điều kiện và yêu cầu hoàn tiền."),
("AIMS Software","Tải đơn hàng, giao dịch thanh toán gốc và lịch sử hoàn tiền."),
("AIMS Software","Xác thực chủ đơn, giao dịch đã thu tiền và điều kiện hoàn."),
("AIMS Software","Tính số tiền còn có thể hoàn, loại trừ khoản đã hoàn hoặc đang xử lý."),
("AIMS Software","Hiển thị đơn, số tiền đủ điều kiện hoàn và thông tin tài khoản nhận theo giao dịch gốc."),
("Customer","Nhập lý do hoàn tiền và kiểm tra số tiền AIMS tính."),
("AIMS Software","Kiểm tra lý do, số tiền và không có yêu cầu hoàn trùng đang chờ."),
("Customer","Xác nhận gửi yêu cầu hoàn tiền."),
("AIMS Software","Gửi yêu cầu tới Interbank kèm mã giao dịch gốc, số tiền và mã chống gửi trùng."),
("Interbank","Xác thực giao dịch gốc và xử lý hoàn về phương thức thanh toán ban đầu."),
("Interbank","Trả kết quả xử lý, mã hoàn tiền và số tiền đã hoàn."),
("AIMS Software","Đối chiếu phản hồi với mã đơn, mã giao dịch gốc và số tiền yêu cầu."),
("AIMS Software","Lưu giao dịch hoàn tiền, thời gian, số tiền, mã hoàn và trạng thái."),
("AIMS Software","Cập nhật đơn/thanh toán sang REFUNDED hoặc REFUND_PENDING theo phản hồi."),
("AIMS Software","Hiển thị và gửi email kết quả hoàn tiền cho Customer.")
],
"alts":[
("3a","AIMS Software","Nếu đơn chưa được thanh toán, không đủ điều kiện hoặc đã hoàn đủ, từ chối yêu cầu và hiển thị lý do; không gửi Interbank."),
("4a","AIMS Software","Nếu còn một yêu cầu hoàn tiền đang xử lý, trả trạng thái hiện tại và không tạo yêu cầu mới trùng lặp."),
("10a","Interbank","Nếu từ chối hoặc timeout, lưu trạng thái FAILED/PENDING theo phản hồi; thông báo Customer và tra cứu trước khi thử lại.")
],
"post":"Nếu thành công, số tiền hoàn và mã giao dịch được lưu, tổng tiền đã hoàn không vượt tổng tiền đã thu và trạng thái đơn/thanh toán được cập nhật. Nếu chưa có xác nhận từ Interbank, giao dịch giữ PENDING/FAILED và không được báo hoàn thành.",
"prompt":"Prompt đã dùng: Hãy đặc tả UC Refund Order của AIMS theo mẫu; giới hạn hoàn theo số tiền đã thu, kiểm tra yêu cầu trùng, gửi Interbank và xử lý phản hồi thành công/từ chối/timeout.",
"eval":"Đánh giá, điều chỉnh: Phản hồi AI ban đầu chưa kiểm soát số dư có thể hoàn và yêu cầu đang chờ. Đặc tả bổ sung đối chiếu giao dịch gốc, giới hạn số tiền, mã chống trùng và tra cứu trạng thái trước khi thử lại để tránh hoàn vượt hoặc hoàn hai lần.",
"io":[
("Mã đơn, mã giao dịch gốc","Vào","AIMS tra cứu từ yêu cầu của Customer hoặc UC004."),
("Lý do yêu cầu hoàn","Vào","Customer cung cấp để lưu dấu vết xử lý."),
("Số tiền còn đủ điều kiện","Ra nội bộ","AIMS tính từ số tiền đã thu trừ số đã hoàn/đang chờ; không vượt khoản đã thu."),
("Yêu cầu hoàn tiền","Ra","Gửi Interbank mã giao dịch gốc, số tiền và mã chống gửi trùng."),
("Kết quả, mã hoàn tiền","Vào từ Interbank","Kết quả thành công/từ chối/chưa xác định, số tiền và mã tham chiếu."),
("Trạng thái và số tiền đã hoàn","Ra","Cập nhật REFUNDED, REFUND_PENDING hoặc FAILED và lưu lịch sử."),
("Thông báo/email","Ra","Nêu số tiền, trạng thái và mã hoàn tiền nếu đã được cấp.")
],
"diagram":{
"nodes":[
("s",0,0,"","start"),("request",1,0,"Chọn đơn, nêu lý do hoàn","action"),
("validate",2,1,"Kiểm tra đã thu tiền và số còn hoàn","action"),("eligible_q",3,1,"Đủ điều kiện, không trùng?","decision"),
("reject",4,0,"Thông báo không đủ điều kiện","end"),("show",4,1,"Hiển thị số tiền có thể hoàn","action"),
("confirm",5,0,"Xác nhận yêu cầu","action"),("send",6,1,"Gửi yêu cầu tới Interbank","action"),
("process",7,2,"Interbank xử lý hoàn tiền","action"),("result_q",8,2,"Hoàn thành?","decision"),
("pending",9,0,"Thông báo chờ/xử lý lỗi","action"),("record",9,1,"Lưu mã hoàn và số tiền","action"),
("update",10,1,"Cập nhật trạng thái đơn","action"),("notify",11,1,"Gửi kết quả cho Customer","action"),
("e",12,0,"Customer nhận kết quả","end")],
"edges":[
("s","request",""),("request","validate",""),("validate","eligible_q",""),
("eligible_q","reject","Không"),("eligible_q","show","Có"),("show","confirm",""),
("confirm","send",""),("send","process",""),("process","result_q",""),
("result_q","pending","Không"),("result_q","record","Có"),("record","update",""),
("update","notify",""),("notify","e","")]
}}
]

def set_cell(cell, value, size=9, bold=False, align=None, color=None):
    p=cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
    p.clear()
    if align is not None: p.alignment=align
    p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.0
    r=p.add_run(str(value))
    r.font.name="Tahoma"; r._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"),"Tahoma")
    r.font.size=Pt(size); r.bold=bold
    if color: r.font.color.rgb=RGBColor(*color)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

def set_main_table(t, steps):
    if len(t.rows)-1 != len(steps): raise ValueError(f"main-flow capacity {len(t.rows)-1}, content {len(steps)}")
    for i,(actor,action) in enumerate(steps,1):
        row=t.rows[i]
        set_cell(row.cells[0],str(i),8.5,align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell(row.cells[1],actor,8.5,align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell(row.cells[2],action,8.5,align=WD_ALIGN_PARAGRAPH.JUSTIFY)

def set_alt_table(t, alts):
    if len(t.rows)-1 != len(alts): raise ValueError(f"alternate-flow capacity {len(t.rows)-1}, content {len(alts)}")
    for i,(number,actor,action) in enumerate(alts,1):
        row=t.rows[i]
        set_cell(row.cells[0],number,8.2,align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell(row.cells[1],actor,8.2,align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell(row.cells[2],action,8.2,align=WD_ALIGN_PARAGRAPH.JUSTIFY)

def imgfont(size,bold=False):
    return ImageFont.truetype(FONT_BOLD if bold else FONT_REG,size)

def wrap_pixels(draw,text,font,max_width):
    words=text.split(); lines=[]; cur=""
    for word in words:
        trial=word if not cur else cur+" "+word
        if draw.textlength(trial,font=font)<=max_width or not cur: cur=trial
        else: lines.append(cur); cur=word
    if cur: lines.append(cur)
    return lines or [""]

def centered_text(draw,box,text,font,fill,spacing=2):
    x0,y0,x1,y1=box
    lines=wrap_pixels(draw,text,font,x1-x0-20)
    bbs=[draw.textbbox((0,0),s,font=font) for s in lines]
    hs=[b[3]-b[1] for b in bbs]; total=sum(hs)+spacing*(len(hs)-1)
    y=y0+(y1-y0-total)/2
    for s,b,h in zip(lines,bbs,hs):
        w=b[2]-b[0]; draw.text((x0+(x1-x0-w)/2,y-b[1]),s,font=font,fill=fill); y+=h+spacing

def draw_diagram(case,path):
    ns={n[0]:{"id":n[0],"row":n[1],"lane":n[2],"text":n[3],"kind":n[4]} for n in case["diagram"]["nodes"]}
    W,lw=1200,400; header,start_y,step=78,126,94
    max_row=max(n["row"] for n in ns.values()); H=header+start_y+max_row*step+115
    im=Image.new("RGB",(W,H),"white"); d=ImageDraw.Draw(im)
    labels=["Customer","AIMS Software","Interbank"]; fills=["#F8FAFC","#F3F7FB","#F8FAFC"]
    hf=imgfont(25,True)
    for lane in range(3):
        x=lane*lw
        d.rectangle((x,0,x+lw,H),fill=fills[lane],outline="#D7DFE8",width=2)
        d.rectangle((x,0,x+lw,header),fill="#E8EEF5",outline="#C8D2DF",width=2)
        centered_text(d,(x+8,0,x+lw-8,header),labels[lane],hf,"#20364D")
    centers=[200,600,1000]; pos={k:(centers[n["lane"]],header+start_y+n["row"]*step) for k,n in ns.items()}
    ef=imgfont(18,True); bh=78
    for src,dst,label in case["diagram"]["edges"]:
        x1,y1=pos[src]; x2,y2=pos[dst]; ks=ns[src]["kind"]; kd=ns[dst]["kind"]
        hs=23 if ks in ("start","end") else (42 if ks=="decision" else bh/2)
        ht=23 if kd in ("start","end") else (42 if kd=="decision" else bh/2)
        sx,sy=x1,y1+hs; ex,ey=x2,y2-ht
        if y2>y1:
            mid=(sy+ey)/2; d.line([(sx,sy),(sx,mid),(ex,mid),(ex,ey)],fill="#526579",width=3,joint="curve")
            d.polygon([(ex,ey),(ex-8,ey-14),(ex+8,ey-14)],fill="#526579")
            if label:
                mx=(sx+ex)/2; ly=mid-23; bb=d.textbbox((0,0),label,font=ef); ww=bb[2]-bb[0]; hh=bb[3]-bb[1]
                d.rounded_rectangle((mx-ww/2-7,ly-3,mx+ww/2+7,ly+hh+5),radius=5,fill="white",outline="#D4DCE5")
                d.text((mx-ww/2,ly),label,font=ef,fill="#34495E")
        else:
            side=38; d.line([(x1-170,y1),(side,y1),(side,y2),(x2-170,y2)],fill="#526579",width=3,joint="curve")
            d.polygon([(x2-170,y2),(x2-184,y2-8),(x2-184,y2+8)],fill="#526579")
            if label: d.text((side+4,(y1+y2)/2),label,font=ef,fill="#34495E")
    af=imgfont(22); df=imgfont(21,True)
    for k,n in ns.items():
        x,y=pos[k]; kind=n["kind"]
        if kind=="start":
            r=23; d.ellipse((x-r,y-r,x+r,y+r),fill="#254A69",outline="#1B354B",width=2)
        elif kind=="end":
            r=23; d.ellipse((x-r,y-r,x+r,y+r),fill="white",outline="#254A69",width=4); d.ellipse((x-r+7,y-r+7,x+r-7,y+r-7),outline="#254A69",width=2)
            if n["text"]: centered_text(d,(x-145,y+25,x+145,y+68),n["text"],imgfont(17,True),"#20364D")
        elif kind=="decision":
            w,h=168,84; pts=[(x,y-h/2),(x+w/2,y),(x,y+h/2),(x-w/2,y)]
            d.polygon(pts,fill="#FFF3D6",outline="#C28A20"); d.line(pts+[pts[0]],fill="#C28A20",width=3)
            centered_text(d,(x-w/2+15,y-h/2+5,x+w/2-15,y+h/2-5),n["text"],df,"#5B4211")
        else:
            w=340; box=(x-w/2,y-bh/2,x+w/2,y+bh/2)
            fill="#F2F8F5" if n["lane"]==2 else "white"; outline="#54816D" if n["lane"]==2 else "#52718F"
            d.rounded_rectangle(box,radius=11,fill=fill,outline=outline,width=3)
            centered_text(d,box,n["text"],af,"#243447")
    im.save(path,optimize=True)

def insert_after(anchor,newp):
    anchor._p.addnext(newp._p)

def add_answers(doc,slot,case):
    anchor=slot
    for text in (case["prompt"],case["eval"]):
        p=doc.add_paragraph(); p.style=doc.styles["Normal"]
        p.paragraph_format.left_indent=Inches(.15); p.paragraph_format.right_indent=Inches(.05)
        p.paragraph_format.space_before=Pt(2); p.paragraph_format.space_after=Pt(4); p.paragraph_format.line_spacing=1.0
        r=p.add_run(text); r.font.name="Tahoma"; r._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"),"Tahoma"); r.font.size=Pt(9.5)
        insert_after(anchor,p); anchor=p

def shade(cell,fill):
    tcpr=cell._tc.get_or_add_tcPr(); shd=tcpr.find(qn("w:shd"))
    if shd is None: shd=OxmlElement("w:shd"); tcpr.append(shd)
    shd.set(qn("w:fill"),fill)

def margins(cell,top=55,start=85,bottom=55,end=85):
    tcpr=cell._tc.get_or_add_tcPr(); mar=tcpr.first_child_found_in("w:tcMar")
    if mar is None: mar=OxmlElement("w:tcMar"); tcpr.append(mar)
    for side,val in (("top",top),("start",start),("bottom",bottom),("end",end)):
        e=mar.find(qn("w:"+side))
        if e is None: e=OxmlElement("w:"+side); mar.append(e)
        e.set(qn("w:w"),str(val)); e.set(qn("w:type"),"dxa")

def add_io_table(doc,slot,rows):
    t=doc.add_table(rows=1,cols=3); t.style="Normal Table"; t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    borders=OxmlElement("w:tblBorders")
    for edge in ("top","left","bottom","right","insideH","insideV"):
        el=OxmlElement("w:"+edge); el.set(qn("w:val"),"single"); el.set(qn("w:sz"),"4"); el.set(qn("w:space"),"0"); el.set(qn("w:color"),"auto"); borders.append(el)
    t._tbl.tblPr.append(borders)
    widths=[Inches(1.62),Inches(1.08),Inches(3.80)]
    for j,text in enumerate(["Dữ liệu","Hướng","Mô tả và điều kiện"]):
        c=t.rows[0].cells[j]; c.width=widths[j]
        set_cell(c,text,8.5,True,WD_ALIGN_PARAGRAPH.CENTER,(32,54,77)); shade(c,"E8EEF5"); margins(c)
    trpr=t.rows[0]._tr.get_or_add_trPr(); repeat=OxmlElement("w:tblHeader"); repeat.set(qn("w:val"),"true"); trpr.append(repeat)
    for row in rows:
        cells=t.add_row().cells
        for j,text in enumerate(row):
            cells[j].width=widths[j]
            set_cell(cells[j],text,8.4,align=WD_ALIGN_PARAGRAPH.CENTER if j==1 else WD_ALIGN_PARAGRAPH.LEFT)
            margins(cells[j])
    slot._p.addnext(t._tbl)

def add_image(doc,slot,path,alt):
    p_el=OxmlElement("w:p"); slot._p.addnext(p_el); p=Paragraph(p_el,slot._parent)
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(3); p.paragraph_format.space_after=Pt(6)
    pic=p.add_run().add_picture(str(path),width=Inches(6.25))
    pic._inline.docPr.set("descr",alt); pic._inline.docPr.set("title",alt)

def stitch(src_path,built_path,out_path):
    with zipfile.ZipFile(src_path) as src, zipfile.ZipFile(built_path) as built:
        src_names=set(src.namelist()); built_names=set(built.namelist())
        relns="http://schemas.openxmlformats.org/package/2006/relationships"
        relroot=etree.fromstring(src.read("word/_rels/document.xml.rels"))
        relbuilt=etree.fromstring(built.read("word/_rels/document.xml.rels"))
        ids={x.get("Id") for x in relroot}
        for rel in relbuilt:
            if rel.get("Id") not in ids: relroot.append(deepcopy(rel))
        relxml=etree.tostring(relroot,xml_declaration=True,encoding="UTF-8",standalone=True)
        ctns="http://schemas.openxmlformats.org/package/2006/content-types"
        ctroot=etree.fromstring(src.read("[Content_Types].xml"))
        if not any(x.tag==f"{{{ctns}}}Default" and x.get("Extension")=="png" for x in ctroot):
            ctroot.insert(0,etree.Element(f"{{{ctns}}}Default",Extension="png",ContentType="image/png"))
        ctxml=etree.tostring(ctroot,xml_declaration=True,encoding="UTF-8",standalone=True)
        wns="http://schemas.openxmlformats.org/wordprocessingml/2006/main"; etree.register_namespace("w",wns)
        settings=etree.fromstring(src.read("word/settings.xml"))
        update=settings.find(f"{{{wns}}}updateFields")
        if update is None: update=etree.SubElement(settings,f"{{{wns}}}updateFields")
        update.set(f"{{{wns}}}val","true")
        settingsxml=etree.tostring(settings,xml_declaration=True,encoding="UTF-8",standalone=True)
        mods={"word/document.xml":built.read("word/document.xml"),"word/_rels/document.xml.rels":relxml,"[Content_Types].xml":ctxml,"word/settings.xml":settingsxml}
        with zipfile.ZipFile(out_path,"w",compression=zipfile.ZIP_DEFLATED) as out:
            for name in src.namelist(): out.writestr(name,mods[name] if name in mods else src.read(name))
            for name in sorted(built_names-src_names):
                if name.startswith("word/media/"): out.writestr(name,built.read(name))

def main():
    doc=Document(str(SRC))
    if len(doc.tables)!=5 or len(doc.sections)!=1: raise ValueError("Template structure differs from expected Bài 2")
    outer=list(doc.tables)
    for i,c in enumerate(cases):
        t=outer[i]
        set_cell(t.rows[0].cells[1],c["code"]); set_cell(t.rows[0].cells[3],c["name"])
        set_cell(t.rows[1].cells[1],c["actors"]); set_cell(t.rows[2].cells[1],c["brief"])
        set_cell(t.rows[3].cells[1],c["pre"]); set_cell(t.rows[6].cells[1],c["post"])
        set_main_table(t.rows[4].cells[1].tables[0],c["steps"])
        set_alt_table(t.rows[5].cells[1].tables[0],c["alts"])
    for p in doc.paragraphs:
        if p.text.startswith("Họ và tên sinh viên:"):
            p.runs[0].text="Họ và tên sinh viên: Bạch Minh Quang"
            for r in p.runs[1:]: r.text=""
        elif p.text.startswith("Mã số sinh viên:"):
            p.runs[0].text="Mã số sinh viên: 202490077"
            for r in p.runs[1:]: r.text=""
    prompts=[p for p in doc.paragraphs if p.text.startswith("<<Nội dung Prompt sử dụng với AI")]
    diagrams=[p for p in doc.paragraphs if p.text.strip()=="<<chèn hình ảnh biểu đồ hoạt động vào đây>>"]
    io_slots=[p for p in doc.paragraphs if p.text.strip()=="<<chèn bảng đặc tả dữ liệu input/output vào đây>>"]
    if not(len(prompts)==len(diagrams)==len(io_slots)==5): raise ValueError("Could not map all five answer slots")
    for c,p in zip(cases,prompts): add_answers(doc,p,c)
    for c,p in zip(cases,diagrams):
        img=IMG_DIR/(c["code"].lower()+"_activity.png"); draw_diagram(c,img); add_image(doc,p,img,"Biểu đồ hoạt động "+c["code"]+" "+c["name"])
    for c,p in zip(cases,io_slots): add_io_table(doc,p,c["io"])
    built=TMP/"python_docx_build.docx"; doc.save(str(built))
    stitch(SRC,built,OUT)
    final=Document(str(OUT))
    print("OUTPUT",OUT); print("SIZE",OUT.stat().st_size)
    print("TOP_LEVEL_TABLES",len(final.tables),"SECTIONS",len(final.sections))
    print("GEOMETRY",final.sections[0].page_width.inches,final.sections[0].page_height.inches)
    print("PROMPTS",len(prompts),"DIAGRAMS",len(diagrams),"IO_TABLES",len(io_slots))
    print("DIAGRAM_FILES")
    for c in cases: print(IMG_DIR/(c["code"].lower()+"_activity.png"))

if __name__=="__main__": main()
