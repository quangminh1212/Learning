from __future__ import annotations

import re
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor


ROOT = Path(r"C:\Dev\Learning\VIII.HUST\Quản trị dự án\Nhóm 1\bao_cao_tong_hop_latex")
SOURCE = ROOT / "main.tex"
ASSETS = ROOT / "assets"
OUTPUT = Path(r"C:\Dev\Learning\tmp\group1_editable_tables.docx")

BLACK = "000000"
LIGHT_BLUE = "EAF1F8"
BORDER = "D9D9D9"
PAGE_WIDTH_MM = 210
LEFT_MM = 30
RIGHT_MM = 22
TEXT_WIDTH_MM = PAGE_WIDTH_MM - LEFT_MM - RIGHT_MM

MACROS = {
    r"\projectname": "Xây dựng hệ thống quản lý hàng hóa và kho bãi sản phẩm thời trang",
    r"\projectcode": "CNTT20261_N01",
}


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size="6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        el = borders.find(tag)
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def set_table_width(table, width_mm: float = TEXT_WIDTH_MM) -> None:
    width_twips = int(width_mm / 25.4 * 1440)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_twips))
    tbl_w.set(qn("w:type"), "dxa")


def set_table_columns(table, widths_mm: list[float]) -> None:
    """Force Word/LibreOffice to honor the calculated native column widths."""
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    grid_cols = list(grid.gridCol_lst)
    while len(grid_cols) < len(widths_mm):
        col = OxmlElement("w:gridCol")
        grid.append(col)
        grid_cols.append(col)
    for index, width_mm in enumerate(widths_mm):
        twips = str(int(width_mm / 25.4 * 1440))
        grid_cols[index].set(qn("w:w"), twips)
        for cell in table.columns[index].cells:
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), twips)
            tc_w.set(qn("w:type"), "dxa")


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_no_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def set_font(run, size=14, bold=False, italic=False, name="Times New Roman") -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), name)


def set_paragraph(p, align=None, before=0, after=4, first_line=0, keep=False) -> None:
    if align is not None:
        p.alignment = align
    fmt = p.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = 1.0
    if first_line:
        fmt.first_line_indent = Cm(first_line)
    fmt.keep_together = keep


def set_style_defaults(doc: Document) -> None:
    for style in doc.styles:
        if style.type != WD_STYLE_TYPE.PARAGRAPH:
            continue
        style.font.name = "Times New Roman"
        style.font.size = Pt(14)
        style.font.color.rgb = RGBColor(0, 0, 0)
        rpr = style._element.get_or_add_rPr()
        rfonts = rpr.rFonts
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.insert(0, rfonts)
        for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
            rfonts.set(qn(f"w:{attr}"), "Times New Roman")

    normal = doc.styles["Normal"]
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.0
    normal.paragraph_format.first_line_indent = Cm(1)

    for name, size, before, after in (
        ("Title", 22, 0, 12),
        ("Heading 1", 19, 14, 8),
        ("Heading 2", 16, 10, 5),
        ("Heading 3", 14, 8, 4),
    ):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.line_spacing = 1.0


def set_section(section, cover=False) -> None:
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    if cover:
        section.top_margin = Mm(12)
        section.bottom_margin = Mm(12)
        section.left_margin = Mm(18)
        section.right_margin = Mm(18)
    else:
        section.top_margin = Mm(25)
        section.bottom_margin = Mm(25)
        section.left_margin = Mm(30)
        section.right_margin = Mm(22)
    section.header_distance = Mm(10)
    section.footer_distance = Mm(10)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, text, fld_end])
    set_font(run, 10)


def add_footer(section) -> None:
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.text = ""
    add_page_field(p)


def strip_comments(text: str) -> str:
    return re.sub(r"(?<!\\)%.*$", "", text)


def read_group(text: str, start: int) -> tuple[str, int]:
    while start < len(text) and text[start].isspace():
        start += 1
    if start >= len(text) or text[start] != "{":
        return "", start
    depth = 0
    out = []
    for i in range(start, len(text)):
        ch = text[i]
        if ch == "{" and (i == 0 or text[i - 1] != "\\"):
            depth += 1
            if depth > 1:
                out.append(ch)
        elif ch == "}" and (i == 0 or text[i - 1] != "\\"):
            depth -= 1
            if depth == 0:
                return "".join(out), i + 1
            out.append(ch)
        else:
            out.append(ch)
    return "".join(out), len(text)


def replace_macros(text: str, refs: dict[str, str] | None = None) -> str:
    refs = refs or {}
    for key, value in MACROS.items():
        text = text.replace(key, value)
    text = re.sub(r"\\vnd\{([^{}]*)\}", r"\1 VNĐ", text)
    text = re.sub(r"\\ref\{([^{}]*)\}", lambda m: refs.get(m.group(1), ""), text)
    text = re.sub(r"\\pageref\{[^{}]*\}", "", text)
    text = re.sub(r"\\(?:tablehead|metric|textbf|textit|emph|texttt|text|mathrm|mathbf|mathit|textsuperscript)\{([^{}]*)\}", r"\1", text)
    text = re.sub(r"\\(?:href|url)\{([^{}]*)\}(?:\{([^{}]*)\})?", lambda m: m.group(2) or m.group(1), text)
    text = re.sub(r"\\begin\{[^{}]*\}|\\end\{[^{}]*\}", "", text)
    text = text.replace(r"\&", "&").replace(r"\_", "_").replace(r"\%", "%")
    text = text.replace(r"\#", "#").replace(r"\$", "$").replace(r"\{", "{").replace(r"\}", "}")
    text = text.replace("~", " ").replace(r"\quad", " ").replace(r"\qquad", " ")
    text = text.replace(r"\,", " ").replace(r"\;", " ").replace(r"\:", " ")
    text = text.replace(r"\linebreak", " ").replace(r"\\", " ")
    text = re.sub(r"\\[a-zA-Z]+(?:\[[^]]*\])?\s*", "", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def plain_cell(text: str, refs: dict[str, str]) -> str:
    text = text.replace("\\\\", "\n")
    text = replace_macros(text, refs)
    return text.strip()


def is_numeric_value(value: str) -> bool:
    return re.fullmatch(r"[0-9][0-9., ]*(?:VNĐ)?", value) is not None


def set_cell_no_wrap(cell) -> None:
    """Prevent Word/LibreOffice from splitting short numeric values at punctuation."""
    tc_pr = cell._tc.get_or_add_tcPr()
    no_wrap = tc_pr.find(qn("w:noWrap"))
    if no_wrap is None:
        no_wrap = OxmlElement("w:noWrap")
        tc_pr.append(no_wrap)


def split_cells(row: str) -> list[str]:
    cells: list[str] = []
    current: list[str] = []
    escaped = False
    depth = 0
    for ch in row:
        if ch == "{" and not escaped:
            depth += 1
        elif ch == "}" and not escaped and depth:
            depth -= 1
        if ch == "&" and not escaped and depth == 0:
            cells.append("".join(current).strip())
            current = []
        else:
            current.append(ch)
        escaped = ch == "\\" and not escaped
        if ch != "\\":
            escaped = False
    cells.append("".join(current).strip())
    return cells


def parse_begin_spec(block: str) -> tuple[str, str]:
    m = re.search(r"\\begin\{(tabularx|longtable|tabular)\}", block)
    if not m:
        return "", ""
    env = m.group(1)
    pos = m.end()
    first, pos = read_group(block, pos)
    if env == "tabularx":
        second, _ = read_group(block, pos)
        return env, second
    return env, first


def parse_widths(spec: str, count: int) -> list[float]:
    widths: list[float] = []
    pos = 0
    while pos < len(spec):
        if spec[pos] in "lrcY":
            widths.append(0.0)
            pos += 1
            continue
        if spec[pos] in "LCY":
            kind = spec[pos]
            pos += 1
            value, end = read_group(spec, pos)
            pos = end
            if kind in "LC" and value.endswith("cm"):
                try:
                    widths.append(float(value[:-2]) * 10.0)
                except ValueError:
                    widths.append(0.0)
            else:
                widths.append(0.0)
            continue
        pos += 1
    if len(widths) != count:
        widths = [0.0] * count
    explicit = sum(widths)
    missing = [i for i, w in enumerate(widths) if w <= 0]
    remaining = max(TEXT_WIDTH_MM - explicit, 0.0)
    if missing:
        each = remaining / len(missing) if remaining else TEXT_WIDTH_MM / count
        widths = [each if w <= 0 else w for w in widths]
    total = sum(widths) or TEXT_WIDTH_MM
    return [w * TEXT_WIDTH_MM / total for w in widths]


def content_fit_widths(widths: list[float], rows: list[list[str]]) -> list[float]:
    mins: list[float] = []
    for col in range(len(widths)):
        values = [row[col].replace("\n", " ").strip() for row in rows if col < len(row)]
        longest = max((len(value) for value in values), default=1)
        nonempty = [value for value in values if value]
        numeric_like = bool(nonempty) and all(
            re.fullmatch(r"[A-Za-z0-9./,%+()\- ]+", value) is not None and len(value) <= 16
            for value in nonempty
        )
        if numeric_like:
            required = max(18.0, min(29.0, longest * 2.65))
        else:
            required = max(18.0, min(55.0, 12.0 + longest * 0.65))
        mins.append(required)
    adjusted = [max(width, minimum) for width, minimum in zip(widths, mins)]
    total = sum(adjusted)
    if total > TEXT_WIDTH_MM:
        factor = TEXT_WIDTH_MM / total
        adjusted = [width * factor for width in adjusted]
    else:
        extra = TEXT_WIDTH_MM - total
        base = sum(widths) or TEXT_WIDTH_MM
        adjusted = [width + extra * original / base for width, original in zip(adjusted, widths)]
    return adjusted


def table_rows(block: str, refs: dict[str, str]) -> tuple[list[list[str]], str]:
    env, spec = parse_begin_spec(block)
    body = block
    begin_line = re.search(r"\\begin\{(?:tabularx|longtable|tabular)\}[^\n]*\n", body)
    if begin_line:
        body = body[begin_line.end() :]
    if env == "longtable" and r"\endfirsthead" in body:
        before, rest = body.split(r"\endfirsthead", 1)
        if r"\endhead" in rest:
            _, after = rest.split(r"\endhead", 1)
            body = before + after
    for marker in (
        r"\endfirsthead",
        r"\endhead",
        r"\endfoot",
        r"\endlastfoot",
        r"\toprule",
        r"\midrule",
        r"\bottomrule",
        r"\hline",
    ):
        body = body.replace(marker, "")
    body = re.sub(r"\\caption\{.*?\}", "", body, flags=re.S)
    body = re.sub(r"\\label\{.*?\}", "", body, flags=re.S)
    rows: list[list[str]] = []
    for raw_row in re.split(r"\\\\(?:\s*\[[^]]*\])?", body):
        raw_row = raw_row.strip()
        if not raw_row or raw_row.startswith(r"\end") or raw_row.startswith(r"\begin"):
            continue
        cells = [plain_cell(c, refs) for c in split_cells(raw_row)]
        if cells and any(c for c in cells):
            rows.append(cells)
    width = max((len(r) for r in rows), default=0)
    for row in rows:
        row.extend([""] * (width - len(row)))
    return rows, spec


def caption_and_label(block: str, refs: dict[str, str]) -> tuple[str, str]:
    m = re.search(r"\\caption\{(.*?)\}", block, flags=re.S)
    caption = plain_cell(m.group(1), refs) if m else ""
    label_m = re.search(r"\\label\{([^{}]*)\}", block)
    label = label_m.group(1) if label_m else ""
    return caption, label


def make_label_map(lines: list[str]) -> tuple[dict[str, str], list[str], list[str]]:
    refs: dict[str, str] = {}
    table_titles: list[str] = []
    figure_titles: list[str] = []
    chapter = 0
    table_no = 0
    figure_no = 0
    mode = ""
    pending: tuple[str, str] | None = None
    for line in lines:
        clean = strip_comments(line).strip()
        m = re.match(r"\\chapter\{(.*?)\}", clean)
        if m:
            chapter += 1
            table_no = 0
            figure_no = 0
        if r"\begin{table}" in clean or r"\begin{longtable}" in clean:
            mode = "table"
        elif r"\begin{figure}" in clean:
            mode = "figure"
        cap = re.search(r"\\caption\{(.*?)\}", clean)
        if cap and mode:
            if mode == "table":
                table_no += 1
                number = f"{chapter}.{table_no}" if chapter else str(table_no)
                title = f"Bảng {number}. {plain_cell(cap.group(1), refs)}"
            else:
                figure_no += 1
                number = f"{chapter}.{figure_no}" if chapter else str(figure_no)
                title = f"Hình {number}. {plain_cell(cap.group(1), refs)}"
            refs[f"__caption__{plain_cell(cap.group(1), refs)}"] = number
            label_m = re.search(r"\\label\{([^{}]*)\}", clean)
            pending = (label_m.group(1), number) if label_m else None
            (table_titles if mode == "table" else figure_titles).append(title)
        label_m = re.search(r"\\label\{([^{}]*)\}", clean)
        if label_m and pending:
            refs[label_m.group(1)] = pending[1]
            pending = None
        if r"\end{table}" in clean or r"\end{longtable}" in clean or r"\end{figure}" in clean:
            mode = ""
    return refs, table_titles, figure_titles


def find_end(lines: list[str], start: int, env: str) -> int:
    end = f"\\end{{{env}}}"
    for i in range(start + 1, len(lines)):
        if end in lines[i]:
            return i
    return len(lines) - 1


def add_text_paragraph(doc: Document, text: str, refs: dict[str, str], style=None, align=None) -> None:
    text = replace_macros(text, refs)
    if not text:
        return
    p = doc.add_paragraph(style=style)
    set_paragraph(p, align=align, first_line=0 if style and style.startswith("Heading") else 1)
    run = p.add_run(text)
    set_font(run, size=14, bold=False)


def add_heading(doc: Document, text: str, level: int, refs: dict[str, str]) -> None:
    text = replace_macros(text, refs)
    p = doc.add_paragraph(style=f"Heading {min(level, 3)}")
    if level == 1 and len(doc.paragraphs) > 2:
        p.paragraph_format.page_break_before = True
    run = p.add_run(text)
    set_font(run, size={1: 19, 2: 16, 3: 14}.get(level, 14), bold=True)


def add_caption(doc: Document, text: str, prefix: str = "") -> None:
    p = doc.add_paragraph()
    set_paragraph(p, before=5, after=4, first_line=0, keep=True)
    run = p.add_run((prefix + " " if prefix else "") + text)
    set_font(run, size=12, bold=True)


def add_native_table(doc: Document, block: str, refs: dict[str, str], caption_prefix: str = "") -> None:
    rows, spec = table_rows(block, refs)
    if not rows:
        return
    caption, label = caption_and_label(block, refs)
    if caption:
        number = refs.get(label, "") or refs.get(f"__caption__{caption}", "")
        prefix = f"Bảng {number}." if number else ""
        add_caption(doc, caption, prefix)
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_width(table)
    set_table_borders(table)
    widths = [TEXT_WIDTH_MM / cols] * cols
    if rows[0][:2] == ["STT", "Tài liệu được sử dụng"] and cols == 2:
        widths = [28.0, 130.0]
    elif caption == "Thành viên dự án" and cols == 3:
        widths = [38.0, 45.0, 75.0]
    elif caption == "Cơ cấu kinh phí dự án" and cols == 3:
        widths = [37.0, 56.0, 65.0]
    elif caption == "Chi phí lao động theo thành viên" and cols == 5:
        widths = [30.0, 29.0, 25.0, 34.0, 40.0]
    elif caption == "Sáu giai đoạn và sản phẩm chính" and cols == 5:
        widths = [24.0, 33.0, 43.0, 28.0, 30.0]
    set_table_columns(table, widths)
    for j, width in enumerate(widths):
        for cell in table.columns[j].cells:
            cell.width = Mm(width)
    for i, row_data in enumerate(rows):
        row = table.rows[i]
        if i == 0:
            repeat_header(row)
        set_no_split(row)
        for j, value in enumerate(row_data):
            cell = row.cells[j]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if caption == "Chi phí lao động theo thành viên" and i == 0:
                set_cell_margins(cell, top=90, start=20, bottom=90, end=20)
            else:
                set_cell_margins(cell)
            if is_numeric_value(value):
                set_cell_no_wrap(cell)
            if i == 0:
                set_cell_shading(cell, LIGHT_BLUE)
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            set_paragraph(p, before=0, after=0, first_line=0)
            r = p.add_run(value)
            set_font(r, size=14, bold=(i == 0 or r"\textbf" in row_data[j]))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_figure(doc: Document, block: str, refs: dict[str, str]) -> None:
    m = re.search(r"\\includegraphics(?:\[([^]]*)\])?\{([^{}]+)\}", block)
    if not m:
        return
    opts = m.group(1) or ""
    asset_name = Path(m.group(2)).name
    asset = ASSETS / asset_name
    if not asset.exists():
        return
    width = TEXT_WIDTH_MM
    wm = re.search(r"width\s*=\s*([0-9.]+)\\textwidth", opts)
    if wm:
        width = TEXT_WIDTH_MM * float(wm.group(1))
    p = doc.add_paragraph()
    set_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER, before=6, after=3, first_line=0, keep=True)
    p.add_run().add_picture(str(asset), width=Mm(width))
    caption, label = caption_and_label(block, refs)
    if caption:
        number = refs.get(label, "") or refs.get(f"__caption__{caption}", "")
        add_caption(doc, caption, f"Hình {number}." if number else "")


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    set_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=8, first_line=0)
    p.add_run().add_picture(str(ASSETS / "logo.png"), width=Mm(34))
    for text, size, gap in (
        ("ĐẠI HỌC BÁCH KHOA HÀ NỘI", 20, 3),
        ("TRƯỜNG CÔNG NGHỆ THÔNG TIN VÀ TRUYỀN THÔNG", 16, 16),
        ("BÁO CÁO TỔNG HỢP DỰ ÁN", 24, 14),
        ("Xây dựng hệ thống quản lý hàng hóa và kho bãi sản phẩm thời trang", 16, 12),
    ):
        p = doc.add_paragraph()
        set_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER, before=gap, after=4, first_line=0, keep=True)
        r = p.add_run(text)
        set_font(r, size=size, bold=True)

    info = [
        ("Mã dự án:", "CNTT20261_N01"),
        ("Thời gian:", "01/09/2026 - 31/03/2027"),
        ("Giám đốc dự án:", "Bùi Tuấn Anh"),
        ("Giảng viên hướng dẫn:", "ThS. Lê Thị Hoa"),
    ]
    table = doc.add_table(rows=len(info), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_width(table, 125)
    set_table_borders(table, color="FFFFFF", size="0")
    set_table_columns(table, [65.0, 60.0])
    for i, (label, value) in enumerate(info):
        table.cell(i, 0).width = Mm(65)
        table.cell(i, 1).width = Mm(60)
        for j, text in enumerate((label, value)):
            cell = table.cell(i, j)
            cell.text = ""
            set_cell_margins(cell, 40, 30, 40, 30)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if j == 0 else WD_ALIGN_PARAGRAPH.LEFT
            set_paragraph(p, before=0, after=0, first_line=0)
            r = p.add_run(text)
            set_font(r, size=14, bold=j == 0)
    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    members = [
        ("Nhóm 1", "Mã sinh viên"),
        ("Bùi Tuấn Anh", "202490032"),
        ("Bùi Quốc Luýt", "202490069"),
        ("Vũ Quang Minh", "202490071"),
        ("Bạch Minh Quang", "202490077"),
        ("Phạm Đoàn Bảo Thiên", "202490090"),
    ]
    table = doc.add_table(rows=len(members), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_width(table, 110)
    set_table_borders(table, color="FFFFFF", size="0")
    set_table_columns(table, [70.0, 40.0])
    for i, (left, right) in enumerate(members):
        for j, text in enumerate((left, right)):
            cell = table.cell(i, j)
            cell.width = Mm(70 if j == 0 else 40)
            cell.text = ""
            set_cell_margins(cell, 20, 20, 20, 20)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            set_paragraph(p, before=0, after=0, first_line=0)
            r = p.add_run(text)
            set_font(r, size=14, bold=i == 0)

    p = doc.add_paragraph()
    set_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER, before=24, after=0, first_line=0)
    r = p.add_run("Hà Nội, tháng 9 năm 2026")
    set_font(r, size=14)


def add_toc(doc: Document, headings: list[tuple[int, str]]) -> None:
    add_heading(doc, "Mục lục", 1, {})
    for level, title in headings:
        p = doc.add_paragraph()
        set_paragraph(p, before=0, after=1, first_line=0)
        p.paragraph_format.left_indent = Cm((level - 1) * 0.6)
        r = p.add_run(title)
        set_font(r, size=14, bold=level == 1)


def parse_heading(line: str) -> tuple[int, str] | None:
    m = re.match(r"\\(chapter|section|subsection|subsubsection)(\*)?\{(.*)\}\s*$", line)
    if not m:
        return None
    return {"chapter": 1, "section": 2, "subsection": 3, "subsubsection": 3}[m.group(1)], m.group(3)


def build() -> None:
    lines = [strip_comments(line.rstrip("\n")) for line in SOURCE.read_text(encoding="utf-8").splitlines()]
    refs, table_titles, figure_titles = make_label_map(lines)
    headings: list[tuple[int, str]] = []
    for line in lines:
        parsed = parse_heading(line.strip())
        if parsed:
            headings.append((parsed[0], replace_macros(parsed[1], refs)))

    doc = Document()
    set_style_defaults(doc)
    set_section(doc.sections[0], cover=True)
    add_cover(doc)
    normal_section = doc.add_section(WD_SECTION.NEW_PAGE)
    set_section(normal_section, cover=False)
    add_footer(normal_section)

    buffer: list[str] = []
    in_body = False
    i = 0

    def flush() -> None:
        if not buffer:
            return
        text = " ".join(x.strip() for x in buffer if x.strip())
        buffer.clear()
        if text:
            add_text_paragraph(doc, text, refs)

    while i < len(lines):
        line = lines[i].strip()
        if line == r"\begin{document}":
            in_body = True
            i += 1
            continue
        if not in_body or line == r"\end{document}":
            i += 1
            continue
        if line == r"\begin{titlepage}":
            i = find_end(lines, i, "titlepage") + 1
            continue
        if line.startswith(r"\begin{table}"):
            flush()
            end = find_end(lines, i, "table")
            add_native_table(doc, "\n".join(lines[i : end + 1]), refs)
            i = end + 1
            continue
        if line.startswith(r"\begin{longtable}"):
            flush()
            end = find_end(lines, i, "longtable")
            add_native_table(doc, "\n".join(lines[i : end + 1]), refs)
            i = end + 1
            continue
        if line.startswith(r"\begin{tabularx}") or line.startswith(r"\begin{tabular}"):
            flush()
            env = "tabularx" if "tabularx" in line else "tabular"
            end = find_end(lines, i, env)
            add_native_table(doc, "\n".join(lines[i : end + 1]), refs)
            i = end + 1
            continue
        if line.startswith(r"\begin{figure}"):
            flush()
            end = find_end(lines, i, "figure")
            add_figure(doc, "\n".join(lines[i : end + 1]), refs)
            i = end + 1
            continue
        if line.startswith(r"\begin{itemize}") or line.startswith(r"\begin{enumerate"):
            flush()
            env = "itemize" if "itemize" in line else "enumerate"
            end = find_end(lines, i, env)
            is_qt_list = env == "enumerate" and r"QT-\arabic*:" in line
            number = 0
            for item_line in lines[i + 1 : end]:
                if r"\item" not in item_line:
                    continue
                item = item_line.split(r"\item", 1)[1].strip()
                if env == "itemize":
                    p = doc.add_paragraph(style="List Bullet")
                    prefix = ""
                else:
                    number += 1
                    p = doc.add_paragraph()
                    p.paragraph_format.left_indent = Cm(0.8)
                    p.paragraph_format.first_line_indent = Cm(-0.8)
                    prefix = f"QT-{number}: " if is_qt_list else f"{number}. "
                set_paragraph(p, before=0, after=2, first_line=0)
                r = p.add_run(prefix + replace_macros(item, refs))
                set_font(r, size=14)
            i = end + 1
            continue
        if line == r"\[":
            flush()
            end = next((j for j in range(i + 1, len(lines)) if lines[j].strip() == r"\]"), i)
            formula = " ".join(lines[i + 1 : end])
            p = doc.add_paragraph()
            set_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER, before=4, after=4, first_line=0)
            r = p.add_run(replace_macros(formula, refs))
            set_font(r, size=14)
            i = end + 1
            continue
        parsed = parse_heading(line)
        if parsed:
            flush()
            add_heading(doc, parsed[1], parsed[0], refs)
            i += 1
            continue
        if line == r"\appendix":
            flush()
            i += 1
            continue
        if line == r"\tableofcontents":
            flush()
            add_toc(doc, headings)
            i += 1
            continue
        if line == r"\listoftables":
            flush()
            add_heading(doc, "Danh sách bảng", 2, refs)
            for title in table_titles:
                p = doc.add_paragraph()
                set_paragraph(p, before=0, after=1, first_line=0)
                r = p.add_run(title)
                set_font(r, size=14)
            i += 1
            continue
        if line == r"\listoffigures":
            flush()
            add_heading(doc, "Danh sách hình", 2, refs)
            for title in figure_titles:
                p = doc.add_paragraph()
                set_paragraph(p, before=0, after=1, first_line=0)
                r = p.add_run(title)
                set_font(r, size=14)
            i += 1
            continue
        if line in (r"\clearpage", r"\newpage"):
            flush()
            p = doc.add_paragraph()
            p.add_run().add_break(WD_BREAK.PAGE)
            i += 1
            continue
        if not line or line.startswith(r"\pagenumbering") or line.startswith(r"\addcontentsline") or line.startswith(r"\label"):
            flush()
            i += 1
            continue
        if line.startswith(r"\vspace") or line in (r"\centering", r"\noindent"):
            i += 1
            continue
        buffer.append(line)
        i += 1
    flush()

    doc.core_properties.title = "Báo cáo tổng hợp dự án Nhóm 1"
    doc.core_properties.subject = "Báo cáo có bảng Word có thể chỉnh sửa"
    doc.core_properties.author = "Nhóm 1"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(f"output={OUTPUT}")
    print(f"headings={len(headings)} tables={len(table_titles)} figures={len(figure_titles)}")


if __name__ == "__main__":
    build()
