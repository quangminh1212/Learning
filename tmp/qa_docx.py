import re
import zipfile
from pathlib import Path

from docx import Document


path = Path(r"C:\Dev\Learning\VIII.HUST\Quản trị dự án\Nhóm 1\bao_cao_tong_hop_latex\output\docx\bao_cao_tong_hop_nhom_1.docx")
doc = Document(path)
xml = zipfile.ZipFile(path).read("word/document.xml").decode("utf-8")
print("paragraphs=", len(doc.paragraphs))
print("tables=", len(doc.tables))
print("inline_shapes=", len(doc.inline_shapes))
print("sections=", len(doc.sections))
print(
    "section_sizes=",
    [
        (
            round(section.page_width.inches, 3),
            round(section.page_height.inches, 3),
            round(section.left_margin.inches, 3),
            round(section.right_margin.inches, 3),
        )
        for section in doc.sections
    ],
)
print("xml_tbl=", xml.count("<w:tbl>"))
print("xml_drawing=", xml.count("<w:drawing>"))
print("direct_colors=", sorted(set(re.findall(r'<w:color[^>]*w:val="([^"]+)"', xml))))
print("font_names=", sorted(set(re.findall(r'<w:rFonts[^>]*w:ascii="([^"]+)"', xml))))
