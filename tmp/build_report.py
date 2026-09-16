"""Build the MS Project screenshot report for Bach Minh Quang."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

ROOT = Path(
    r"C:\Dev\Learning\VIII.HUST\Quản trị dự án\BachMinhQuang_202490077"
)
SHOT = ROOT / "QLDA_BTCN_202490077_BachMinhQuang" / "screenshots"
OUT = ROOT / "BaocaoTongHop_202490077_BachMinhQuang.docx"

BAI_INFO = {
    1: ("9 ngày", "A-B-D-H-I"),
    2: ("24 ngày", "B-E-F-H-I-J"),
    3: ("42 ngày", "A-D-G-F-L-P và A-D-G-F-K-N-P"),
    4: ("37 ngày", "A-C-G-H-I-K-N"),
    5: ("21 ngày", "B-D-H-K"),
}

VIEWS = (
    ("5ngay", "gantt", "2.1 Tuần làm việc 5 ngày — Gantt (bảng công việc)"),
    ("5ngay", "network", "2.1 Tuần làm việc 5 ngày — Network Diagram (sơ đồ găng)"),
    ("7ngay", "gantt", "2.2 Tuần làm việc 7 ngày — Gantt (bảng công việc)"),
    ("7ngay", "network", "2.2 Tuần làm việc 7 ngày — Network Diagram (sơ đồ găng)"),
)


def set_run_font(run, name="Times New Roman", size=12, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_heading(doc, text, size=16):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=True, color=RGBColor(0x1F, 0x4E, 0x79))
    return p


def add_body(doc, text, size=12):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_run_font(run, size=size)
    return p


def add_picture_page(doc, image_path: Path, caption: str, filename: str) -> None:
    doc.add_page_break()
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = cap.add_run(caption)
    set_run_font(run, size=13, bold=True, color=RGBColor(0x1F, 0x4E, 0x79))
    fn = doc.add_paragraph()
    run = fn.add_run(f"Tên file: {filename}")
    set_run_font(run, size=10, color=RGBColor(0x55, 0x55, 0x55))
    fn.paragraph_format.space_after = Pt(6)

    section = doc.sections[-1]
    usable_w = section.page_width - section.left_margin - section.right_margin
    usable_h = section.page_height - section.top_margin - section.bottom_margin - Cm(2.2)
    with Image.open(image_path) as im:
        w_px, h_px = im.size
    aspect = h_px / w_px
    width = usable_w
    height = int(width * aspect)
    if height > usable_h:
        height = usable_h
        width = int(height / aspect)
    doc.add_picture(str(image_path), width=width)


def main() -> None:
    missing = []
    for bai in range(1, 6):
        for tag, view, _ in VIEWS:
            name = f"btcn{bai}_{tag}_{view}_202490077_BachMinhQuang.png"
            if not (SHOT / name).exists():
                missing.append(name)
    if missing:
        raise SystemExit("missing shots: " + ", ".join(missing))

    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.left_margin = Cm(1.2)
    section.right_margin = Cm(1.2)
    section.top_margin = Cm(1.2)
    section.bottom_margin = Cm(1.2)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = header.add_run("202490077 - Bạch Minh Quang")
    set_run_font(r, size=11, bold=True)
    r2 = header.add_run("\tBáo cáo thực hành cá nhân")
    set_run_font(r2, size=11, bold=True)
    header.paragraph_format.tab_stops.add_tab_stop(Cm(27.0), alignment=2)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run("Quản trị dự án — MS Project — BTCN 1–5")
    set_run_font(fr, size=9, color=RGBColor(0x66, 0x66, 0x66))

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Cm(3.5)
    run = title.add_run("ĐẠI HỌC BÁCH KHOA HÀ NỘI")
    set_run_font(run, size=16, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Trường Công nghệ Thông tin và Truyền thông")
    set_run_font(run, size=13, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Cm(1.4)
    run = p.add_run("BÁO CÁO THỰC HÀNH CÁ NHÂN")
    set_run_font(run, size=22, bold=True, color=RGBColor(0x1F, 0x4E, 0x79))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("QUẢN TRỊ DỰ ÁN — MICROSOFT PROJECT")
    set_run_font(run, size=16, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Cm(1.6)
    run = p.add_run("Sinh viên: Bạch Minh Quang")
    set_run_font(run, size=14)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("MSSV: 202490077")
    set_run_font(run, size=14)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Hà Nội, 09/2026")
    set_run_font(run, size=13)

    doc.add_page_break()
    add_heading(doc, "Nội dung báo cáo", 18)
    add_body(
        doc,
        "Mỗi bài tập gồm file MS Project lịch 5 ngày và lịch 7 ngày. "
        "Mỗi lịch chụp 2 ảnh: Gantt (bảng công việc) và Network Diagram (sơ đồ găng). "
        "Tổng 5 bài × 4 ảnh = 20 ảnh. Phần tính PERT/CPM trình bày trong báo cáo LaTeX.",
    )
    add_body(doc, "Quy ước tên file ảnh: btcn{N}_{5ngay|7ngay}_{gantt|network}_202490077_BachMinhQuang.png")
    add_body(doc, "Quy ước tên file MPP: btcn{N}_{5ngay|7ngay}_202490077_BachMinhQuang.mpp")
    add_body(doc, "Ngày bắt đầu dự án: 15/09/2026 08:00. Tên công việc: {ký hiệu}_0077_BachMinhQuang.")

    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(["Bài", "Thời gian", "Đường găng", "File MPP"]):
        hdr[i].text = h
        for p in hdr[i].paragraphs:
            for run in p.runs:
                set_run_font(run, size=11, bold=True)
    for bai, (dur, path) in BAI_INFO.items():
        row = table.add_row().cells
        vals = [
            f"Bài {bai}",
            dur,
            path,
            f"btcn{bai}_5ngay_…mpp\nbtcn{bai}_7ngay_…mpp",
        ]
        for i, v in enumerate(vals):
            row[i].text = v
            for p in row[i].paragraphs:
                for run in p.runs:
                    set_run_font(run, size=11)

    for bai, (dur, path) in BAI_INFO.items():
        doc.add_page_break()
        add_heading(doc, f"Bài tập {bai}", 18)
        add_body(doc, f"Thời gian dự án: {dur}. Đường găng: {path}.")
        add_body(
            doc,
            "2. Tạo dự án trên MS Project (bảng công việc, sơ đồ găng). "
            "Mỗi mục dưới đây là một ảnh chụp cửa sổ Microsoft Project, "
            "đã ghi rõ bài / lịch / kiểu view trên đầu ảnh và trong tên file.",
        )
        for tag, view, caption in VIEWS:
            fname = f"btcn{bai}_{tag}_{view}_202490077_BachMinhQuang.png"
            add_picture_page(doc, SHOT / fname, f"Bài {bai} — {caption}", fname)

    doc.save(str(OUT))
    print("wrote", OUT, OUT.stat().st_size)


if __name__ == "__main__":
    main()
