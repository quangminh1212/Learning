from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile
import shutil
import tempfile
import xml.etree.ElementTree as ET

from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from pdf2docx import Converter


PDF_PATH = Path(
    r"C:\Dev\Learning\VIII.HUST\Quản trị dự án\Nhóm 1\bao_cao_tong_hop_latex\output\pdf\bao_cao_tong_hop_nhom_1.pdf"
)
RAW_PATH = Path(r"C:\Dev\Learning\tmp\group1_pdf2docx_raw.docx")
FINAL_PATH = Path(
    r"C:\Dev\Learning\VIII.HUST\Quản trị dự án\Nhóm 1\bao_cao_tong_hop_latex\output\docx\bao_cao_tong_hop_nhom_1_editable.docx"
)


def convert_pdf() -> None:
    converter = Converter(str(PDF_PATH))
    try:
        converter.convert(str(RAW_PATH), start=0, end=None)
    finally:
        converter.close()


def set_run_style(run) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        from docx.oxml import OxmlElement

        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{key}"), "Times New Roman")


def set_every_run_font(doc_path: Path) -> None:
    doc = Document(doc_path)
    for style in doc.styles:
        try:
            style.font.name = "Times New Roman"
            style.font.size = Pt(14)
            style.font.color.rgb = RGBColor(0, 0, 0)
        except (AttributeError, ValueError):
            pass
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            set_run_style(run)
    def format_table(table) -> None:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        set_run_style(run)
                for nested_table in cell.tables:
                    format_table(nested_table)

    for table in doc.tables:
        format_table(table)
    for section in doc.sections:
        for container in (section.header, section.footer):
            for paragraph in container.paragraphs:
                for run in paragraph.runs:
                    set_run_style(run)
            for table in container.tables:
                format_table(table)

    for section in doc.sections:
        section.page_width = doc.sections[0].page_width
        section.page_height = doc.sections[0].page_height
    doc.core_properties.title = "Báo cáo tổng hợp dự án Nhóm 1"
    doc.core_properties.subject = "Bản Word có thể chỉnh sửa"
    doc.core_properties.author = "Nhóm 1"
    doc.save(doc_path)


def verify_xml_font(doc_path: Path) -> tuple[int, int, int]:
    with ZipFile(doc_path) as zf:
        xml = zf.read("word/document.xml")
    root = ET.fromstring(xml)
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    runs = root.findall(".//w:r", ns)
    tnr = 0
    black = 0
    size14 = 0
    for run in runs:
        rpr = run.find("w:rPr", ns)
        if rpr is None:
            continue
        fonts = rpr.find("w:rFonts", ns)
        if fonts is not None and fonts.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}ascii") == "Times New Roman":
            tnr += 1
        color = rpr.find("w:color", ns)
        if color is not None and color.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val") == "000000":
            black += 1
        size = rpr.find("w:sz", ns)
        if size is not None and size.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val") == "28":
            size14 += 1
    return len(runs), tnr, black + size14


if __name__ == "__main__":
    FINAL_PATH.parent.mkdir(parents=True, exist_ok=True)
    convert_pdf()
    set_every_run_font(RAW_PATH)
    shutil.copy2(RAW_PATH, FINAL_PATH)
    total, tnr, black_or_size = verify_xml_font(FINAL_PATH)
    print(f"raw_runs={total} tnr_runs={tnr} black_or_size_checks={black_or_size}")
    print("output=" + str(FINAL_PATH).encode("unicode_escape").decode("ascii"))
