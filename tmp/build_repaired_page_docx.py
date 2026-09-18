from pathlib import Path

import fitz
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor


PDF_PATH = Path(r"C:\Dev\Learning\tmp\group1_latex_black_1\build\main.pdf")
PNG_DIR = Path(r"C:\Dev\Learning\tmp\group1_docx_pages_black")
DOCX_PATH = Path(r"C:\Dev\Learning\tmp\group1_docx_repaired.docx")


def render_pages() -> list[Path]:
    PNG_DIR.mkdir(parents=True, exist_ok=True)
    pdf = fitz.open(PDF_PATH)
    image_paths: list[Path] = []
    scale = 300 / 72
    for index, page in enumerate(pdf, start=1):
        path = PNG_DIR / f"page-{index:02d}.png"
        pixmap = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
        pixmap.save(path)
        image_paths.append(path)
    pdf.close()
    return image_paths


def apply_word_defaults(document: Document) -> None:
    for style in document.styles:
        if style.type != WD_STYLE_TYPE.PARAGRAPH:
            continue
        style.font.name = "Times New Roman"
        style.font.size = Pt(14)
        style.font.color.rgb = RGBColor(0, 0, 0)
        rpr = style._element.get_or_add_rPr()
        rfonts = rpr.rFonts
        if rfonts is None:
            from docx.oxml import OxmlElement

            rfonts = OxmlElement("w:rFonts")
            rpr.insert(0, rfonts)
        for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
            rfonts.set(qn(f"w:{attr}"), "Times New Roman")


def build_docx(image_paths: list[Path]) -> None:
    document = Document()
    apply_word_defaults(document)

    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(0)
    section.bottom_margin = Mm(0)
    section.left_margin = Mm(0)
    section.right_margin = Mm(0)
    section.header_distance = Mm(0)
    section.footer_distance = Mm(0)

    document.core_properties.title = "Báo cáo tổng hợp dự án Nhóm 1"
    document.core_properties.subject = "DOCX giữ nguyên trình bày theo tài liệu PDF"
    document.core_properties.author = "Nhóm 1"

    for index, image_path in enumerate(image_paths):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.left_indent = Pt(0)
        paragraph.paragraph_format.right_indent = Pt(0)
        paragraph.paragraph_format.line_spacing = 1
        paragraph.paragraph_format.keep_together = True
        paragraph.paragraph_format.widow_control = False
        paragraph.paragraph_format.page_break_before = index > 0
        run = paragraph.add_run()
        run.font.name = "Times New Roman"
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(0, 0, 0)
        run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
        run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
        run.add_picture(str(image_path), width=Mm(209.8), height=Mm(296.8))

    document.save(DOCX_PATH)


if __name__ == "__main__":
    pages = render_pages()
    build_docx(pages)
    print(f"created_pages={len(pages)}")
    print(f"output={DOCX_PATH}")
