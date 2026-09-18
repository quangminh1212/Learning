from pathlib import Path

import fitz
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Mm, Pt


PDF_PATH = Path(
    r"C:\Dev\Learning\VIII.HUST\Quản trị dự án\Nhóm 1\bao_cao_tong_hop_latex\output\pdf\bao_cao_tong_hop_nhom_1.pdf"
)
PNG_DIR = Path(r"C:\Dev\Learning\tmp\group1_docx_pages")
DOCX_PATH = Path(
    r"C:\Dev\Learning\VIII.HUST\Quản trị dự án\Nhóm 1\bao_cao_tong_hop_latex\output\docx\bao_cao_tong_hop_nhom_1.docx"
)


def render_pages() -> list[Path]:
    PNG_DIR.mkdir(parents=True, exist_ok=True)
    pdf = fitz.open(PDF_PATH)
    image_paths: list[Path] = []
    scale = 300 / 72
    for index, page in enumerate(pdf, start=1):
        path = PNG_DIR / f"page-{index:02d}.png"
        pixmap = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
        pixmap.save(path)
        image_paths.append(path)
    pdf.close()
    return image_paths


def build_docx(image_paths: list[Path]) -> None:
    DOCX_PATH.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(0)
    section.bottom_margin = Mm(0)
    section.left_margin = Mm(0)
    section.right_margin = Mm(0)
    section.header_distance = Mm(0)
    section.footer_distance = Mm(0)

    document.core_properties.title = "Báo cáo tổng hợp dự án Nhóm 1"
    document.core_properties.subject = "Bản DOCX giữ nguyên trình bày theo PDF"
    document.core_properties.author = "Nhóm 1"

    for index, image_path in enumerate(image_paths):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.left_indent = Pt(0)
        paragraph.paragraph_format.right_indent = Pt(0)
        paragraph.paragraph_format.line_spacing = 1
        paragraph.paragraph_format.page_break_before = index > 0
        run = paragraph.add_run()
        run.add_picture(str(image_path), width=Mm(209.8), height=Mm(296.8))

    document.save(DOCX_PATH)


if __name__ == "__main__":
    pages = render_pages()
    build_docx(pages)
    print(f"created_pages={len(pages)}")
    print("output=" + str(DOCX_PATH).encode("unicode_escape").decode("ascii"))
